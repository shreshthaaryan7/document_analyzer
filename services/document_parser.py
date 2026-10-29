from pathlib import Path
import io

def extract_document(data: bytes, extension: str):
    ext=extension.lower().lstrip('.')
    if ext=='txt':
        text=data.decode('utf-8-sig',errors='replace').strip()
        return ([{'text':text,'source_ref':'Text'}] if text else []), None
    if ext=='pdf':
        import fitz
        pages=[]
        with fitz.open(stream=data,filetype='pdf') as pdf:
            if pdf.is_encrypted and not pdf.authenticate(''): raise ValueError('This PDF is password-protected.')
            for i,page in enumerate(pdf):
                t=page.get_text('text').strip()
                if t: pages.append({'text':t,'source_ref':f'Page {i+1}'})
            count=len(pdf)
        if not pages: raise ValueError('No selectable text was found. Scanned PDFs may require OCR, which is not enabled.')
        return pages,count
    if ext=='docx':
        from docx import Document
        doc=Document(io.BytesIO(data)); parts=[]
        for p in doc.paragraphs:
            if p.text.strip(): parts.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows: parts.append(' | '.join(c.text.strip() for c in row.cells))
        text='\n'.join(parts).strip()
        if not text: raise ValueError('No readable text was found in this DOCX file.')
        return [{'text':text,'source_ref':'Document section'}],None
    raise ValueError('Unsupported file type. Upload PDF, DOCX, or TXT.')
